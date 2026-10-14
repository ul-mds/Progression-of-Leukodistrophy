#!/usr/bin/env python3
"""Build Journal of Neurology cover letter with logos and author expertise."""
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = Path(__file__).resolve().parent
LOGOS = sorted(HERE.glob("cover_logos/cover_logo_*.png"))
OUT = HERE / "Cover_letter.docx"


def add_para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic
    return p


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    if len(LOGOS) >= 4:
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, logo in enumerate(LOGOS[:4]):
            cell = table.rows[0].cells[i]
            p = cell.paragraphs[0]
            p.alignment = WD_TABLE_ALIGNMENT.CENTER
            run = p.add_run()
            run.add_picture(str(logo), width=Inches(1.35))
        doc.add_paragraph()

    add_para(doc, "14 October 2026")
    add_para(doc, "")
    add_para(doc, "Editor-in-Chief")
    add_para(doc, "Journal of Neurology")
    add_para(doc, "")
    add_para(doc, "Dear Editor-in-Chief,")
    add_para(doc, "")
    add_para(
        doc,
        "Please consider the enclosed manuscript for publication as an Original Communication in the Journal of Neurology.",
    )
    add_para(
        doc,
        "Title: Progression from myeloneuropathy to cerebral disease in men with X-linked adrenoleukodystrophy studied using a Human Phenotype Ontology-based approach",
        bold=True,
    )
    add_para(
        doc,
        "The manuscript has not been published elsewhere and is not under consideration by another journal. All authors have read and approved the final version and agree with its submission.",
    )
    add_para(doc, "")
    add_para(
        doc,
        "Men with X-linked adrenoleukodystrophy (X-ALD) may remain on a myeloneuropathic trajectory for years before inflammatory cerebral adrenoleukodystrophy (cerebral ALD) is recognised. In clinical practice, estimates of when and how often this transition occurs remain inconsistent, which complicates counselling and the design of surveillance and natural history studies. We address this question using longitudinal data from the Leipzig leukodystrophy registry (REDCap), a setting in which neurologists record disease manifestation and Human Phenotype Ontology (HPO)-coded examination findings at repeated visits.",
    )
    add_para(
        doc,
        "In a retrospective cohort of 146 men with a myelopathic onset, 65 had cerebral ALD at the first dated visit. Among 67 men without cerebral ALD at entry and with follow-up (median 86 months; 20 conversion events), the cumulative incidence of a first physician-recorded cerebral ALD event was 6.5% at 3 years, 10.2% at 5 years, and 24.3% (95% CI 13.8 to 40.7) at 10 years. Baseline clinical and HPO-coded features did not identify converters after adjustment for multiple testing. We relate these estimates to earlier cohorts and emphasise that the analyses are exploratory; magnetic resonance imaging surveillance remains the standard for detecting cerebral disease.",
    )
    add_para(
        doc,
        "A central element of the paper is an original HPO Symptom Grouping Dataset (Online Resource 1), developed by the clinical authors for this study. It maps 135 registry examination items to 133 distinct HPO terms in myelopathic and cerebral subgroups with neuroanatomical localisation. To our knowledge, standardised HPO-coded phenotypes have not previously been used to study progression from myeloneuropathy to cerebral disease in leukodystrophy. The dataset and submission materials are available at https://github.com/ul-mds/Progression-of-Leukodistrophy.",
    )
    add_para(
        doc,
        "The clinical mapping reflects phenotyping practice at our centre. Professor Wolfgang Köhler, who co-developed the dataset, has long-standing expertise in the diagnosis and longitudinal care of leukodystrophy, including early clinical characterisation of these disorders in adult neurology. The study is reported according to the STROBE and RECORD statements; funding, competing interests, data availability, and the distinction between this analysis and an earlier exploratory registry exercise are documented in the manuscript.",
    )
    add_para(doc, "")
    add_para(
        doc,
        "We respectfully suggest that the work will be of interest to clinicians and researchers concerned with X-ALD, leukodystrophy registries, and interoperable clinical phenotype data. We would be grateful if you could arrange peer review. Institutional e-mail addresses for suggested reviewers can be provided on request.",
    )
    add_para(doc, "")
    add_para(doc, "Thank you for your consideration.")
    add_para(doc, "")
    add_para(doc, "Yours sincerely,")
    add_para(doc, "")
    add_para(doc, "Senol Dogan, PhD", bold=True)
    add_para(doc, "Corresponding author")
    add_para(doc, "Institute for Medical Informatics, Statistics and Epidemiology")
    add_para(doc, "Medical Faculty, Leipzig University")
    add_para(doc, "Härtelstr. 16–18, 04107 Leipzig, Germany")
    add_para(doc, "Senol.Dogan@medizin.uni-leipzig.de")
    add_para(doc, "ORCID: 0000-0001-7358-4850")
    add_para(doc, "")
    add_para(doc, "On behalf of all co-authors:")
    add_para(doc, "Wolfgang Köhler, Caroline Bergner, Julia Lier, Sina Sadeghi, Toralf Kirsten")
    add_para(doc, "")
    add_para(doc, "Relevant experience", bold=True)
    expertise = [
        "Senol Dogan. He is a bioinformatician and biostatistician in the Department for Medical Data Science at Leipzig University, with a doctorate in genetics and bioengineering. For this manuscript he designed the analysis, performed the statistical work, and wrote the first draft. He is the corresponding author.",
        "Toralf Kirsten. He is Professor of Medical Data Science at Leipzig University. His research concerns the analysis of clinical registry data in rare neurological disease. He supervised the project and revised the manuscript.",
        "Wolfgang Köhler. He is a neurologist at Leipzig University Medical Center and head of the leukodystrophy clinic at the Myelin Centre. He cares for adults with leukodystrophy, co-developed the HPO Symptom Grouping Dataset for this study, and contributed long-standing clinical expertise in phenotyping and natural history. He was among the physicians involved in the early clinical recognition of leukodystrophy in adult neurology.",
        "Caroline Bergner. She is a neurologist at Leipzig University Medical Center. She leads the Leukodystrophy Laboratory and the adult leukodystrophy outpatient clinic. Her research concerns how hereditary leukodystrophies present in adults. She helped to interpret the data and revised the manuscript.",
        "Julia Lier. She is a clinician scientist in the Leukodystrophy Laboratory and a physician in the Department of Neurology at Leipzig. She works with adults who have leukodystrophy and contributed to the clinical interpretation of the registry phenotypes used here.",
        "Sina Sadeghi. She is a researcher in medical data science at Leipzig University. She contributed to the analysis and revised the manuscript.",
    ]
    for block in expertise:
        add_para(doc, block)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
